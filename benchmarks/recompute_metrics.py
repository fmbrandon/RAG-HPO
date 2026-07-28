#!/usr/bin/env python3
"""Recompute precision, recall, and F1 from workbook TP/FP/FN counts."""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

COMPARISON_SHEETS = (
    "Premium LLM Comparison",
    "CSC Comparison Results",
    "GSC Comparison Results",
)


@dataclass(frozen=True)
class MetricRow:
    sheet: str
    model: str
    patient_id: str
    tp: int
    fp: int
    fn: int
    precision: float
    recall: float
    f1: float


def ratio(numerator: int | float, denominator: int | float) -> float:
    return numerator / denominator if denominator else 0.0


def compute(tp: int, fp: int, fn: int) -> tuple[float, float, float]:
    precision = ratio(tp, tp + fp)
    recall = ratio(tp, tp + fn)
    f1 = ratio(2 * precision * recall, precision + recall)
    return precision, recall, f1


def extract(workbook: Path) -> list[MetricRow]:
    book = load_workbook(workbook, read_only=True, data_only=True)
    rows: list[MetricRow] = []
    for sheet_name in COMPARISON_SHEETS:
        sheet = book[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
        if len(values) < 3:
            continue
        for start in range(0, len(values[0]), 8):
            model = values[0][start + 1] if start + 1 < len(values[0]) else None
            if not model:
                continue
            for value_row in values[2:]:
                if start + 3 >= len(value_row):
                    continue
                patient_id = value_row[start]
                tp = value_row[start + 1]
                fp = value_row[start + 2]
                fn = value_row[start + 3]
                if patient_id is None or not all(
                    isinstance(value, (int, float)) for value in (tp, fp, fn)
                ):
                    continue
                precision, recall, f1 = compute(int(tp), int(fp), int(fn))
                rows.append(
                    MetricRow(
                        sheet=sheet_name,
                        model=str(model).strip(),
                        patient_id=str(patient_id),
                        tp=int(tp),
                        fp=int(fp),
                        fn=int(fn),
                        precision=precision,
                        recall=recall,
                        f1=f1,
                    )
                )
    return rows


def write(rows: list[MetricRow], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(MetricRow.__annotations__))
        writer.writeheader()
        writer.writerows(row.__dict__ for row in rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    rows = extract(args.workbook)
    if not rows:
        parser.error("no comparison rows were found")
    write(rows, args.output)
    print(f"Wrote {len(rows)} metric rows to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
